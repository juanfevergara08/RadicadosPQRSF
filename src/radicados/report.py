"""Generación del reporte diario en Excel con los correos no radicados.

El Excel no se calcula aparte: es una vista de las filas en estado FALLIDO o
REVISION de la base de estado. Así el reporte nunca contradice al proceso.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import EstadoProceso
from .state import AlmacenEstado

CABECERAS = [
    "Message-ID",
    "Fecha del correo",
    "Remitente",
    "Asunto",
    "Estado",
    "Motivo / error",
    "Intentos",
    "Última actualización",
]

ANCHOS = [38, 20, 34, 46, 12, 44, 10, 22]

ENCABEZADO_FONDO = PatternFill("solid", fgColor="1F4E78")
ENCABEZADO_FUENTE = Font(color="FFFFFF", bold=True)


def generar_reporte(
    almacen: AlmacenEstado,
    directorio_salida: str | Path,
    fecha: date | None = None,
) -> Path:
    """Escribe el .xlsx y devuelve su ruta."""
    fecha = fecha or date.today()
    salida = Path(directorio_salida)
    salida.mkdir(parents=True, exist_ok=True)
    ruta = salida / f"radicados_no_procesados_{fecha.isoformat()}.xlsx"

    fallidos = almacen.listar_por_estado(EstadoProceso.FALLIDO)
    revision = almacen.listar_por_estado(EstadoProceso.REVISION)

    libro = Workbook()

    hoja = libro.active
    hoja.title = "No radicados"

    hoja.append(CABECERAS)
    for celda in hoja[1]:
        celda.fill = ENCABEZADO_FONDO
        celda.font = ENCABEZADO_FUENTE
        celda.alignment = Alignment(vertical="center")

    for fila in fallidos + revision:
        motivo = (
            fila["ultimo_error"]
            if fila["estado"] == EstadoProceso.FALLIDO.value
            else fila["motivo_revision"]
        )
        hoja.append(
            [
                fila["message_id"],
                fila["recibido_en"] or "",
                fila["remitente"],
                fila["asunto"],
                fila["estado"],
                motivo,
                fila["intentos"],
                fila["actualizado_en"],
            ]
        )

    for indice, ancho in enumerate(ANCHOS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = "A2"
    if hoja.max_row > 1:
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(CABECERAS))}{hoja.max_row}"

    resumen = libro.create_sheet("Resumen")
    resumen.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    resumen.append([])
    resumen.append(["Estado", "Cantidad"])
    for celda in resumen[3]:
        celda.font = Font(bold=True)

    conteo = almacen.resumen()
    for estado in EstadoProceso:
        resumen.append([estado.value, conteo.get(estado.value, 0)])
    resumen.append(["total", sum(conteo.values())])
    resumen["A1"].font = Font(bold=True)
    resumen.column_dimensions["A"].width = 18
    resumen.column_dimensions["B"].width = 14

    libro.save(ruta)
    return ruta
